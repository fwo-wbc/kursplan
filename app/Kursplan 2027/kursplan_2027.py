import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date, timedelta

def generate_2027_schedule_dynamic_mix():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Terminplanung 2027"
    ws.views.sheetView[0].showGridLines = True

    # 1. Feiertage 2027
    feiertage = {
        date(2027, 1, 1): "Neujahr",
        date(2027, 3, 26): "Karfreitag",
        date(2027, 3, 29): "Ostermontag",
        date(2027, 5, 1): "Tag der Arbeit",
        date(2027, 5, 6): "Christi Himmelfahrt",
        date(2027, 5, 17): "Pfingstmontag",
        date(2027, 5, 27): "Fronleichnam",
        date(2027, 10, 3): "Tag der Dt. Einheit",
        date(2027, 10, 31): "Reformationstag",
        date(2027, 11, 1): "Allerheiligen",
        date(2027, 12, 25): "1. Weihnachtstag",
        date(2027, 12, 26): "2. Weihnachtstag",
    }

    # Brückentage (Do -> Fr frei, Di -> Mo frei)
    brueckentage = {
        date(2027, 5, 7): "Brückentag (nach Himmelfahrt)",
        date(2027, 5, 28): "Brückentag (nach Fronleichnam)",
    }

    # Urlaubszeiten (KW 1, KW 16, KW 29-33, KW 51-01)
    urlaub_ranges = [
        (date(2027, 1, 1), date(2027, 1, 8)),
        (date(2027, 4, 19), date(2027, 4, 25)),
        (date(2027, 7, 19), date(2027, 8, 22)),
        (date(2027, 12, 20), date(2028, 1, 9)),
    ]

    def is_urlaub(d):
        for s, e in urlaub_ranges:
            if s <= d <= e:
                return True
        return False

    def is_blocked(d):
        return (d in feiertage) or (d in brueckentage) or is_urlaub(d)

    # 2. Rotierende Prio-3-Permutationen für 11 Zyklen (ohne Software-Doppelung an einem Tag)
    # Slots:
    # 0: Partner PowerPoint Grundlagen (P1)
    # 1: Partner Word Grundlagen (P2)
    # 2: Partner Excel Grundlagen Modul 2 (P1)
    # 3: Partner PowerPoint Aufbau (P1)
    # 4: Partner Word Aufbau (P2)
    # 5: Partner Excel Grundlagen Modul 3 (P2)
    prio3_rotations = [
        # Zyklus 1
        [
            'Word: Seriendruck',
            'PowerPoint: Präsentation vorbereiten',
            'Word: Dokumentvorlagen',
            'Word: Umfangreiche Dokumente',
            'Excel: Power Query Grundlagen',
            'PowerPoint: Designvorlagen'
        ],
        # Zyklus 2
        [
            'Word: Dokumentvorlagen',
            'Excel: Power Query Grundlagen',
            'Word: Umfangreiche Dokumente',
            'Word: Seriendruck',
            'PowerPoint: Designvorlagen',
            'PowerPoint: Präsentation vorbereiten'
        ],
        # Zyklus 3
        [
            'Word: Dokumentvorlagen',
            'PowerPoint: Designvorlagen',
            'PowerPoint: Präsentation vorbereiten',
            'Word: Seriendruck',
            'Excel: Power Query Grundlagen',
            'Word: Umfangreiche Dokumente'
        ],
        # Zyklus 4
        [
            'Word: Seriendruck',
            'Excel: Power Query Grundlagen',
            'PowerPoint: Designvorlagen',
            'Word: Umfangreiche Dokumente',
            'PowerPoint: Präsentation vorbereiten',
            'Word: Dokumentvorlagen'
        ],
        # Zyklus 5
        [
            'Excel: Power Query Grundlagen',
            'PowerPoint: Designvorlagen',
            'Word: Umfangreiche Dokumente',
            'Word: Dokumentvorlagen',
            'PowerPoint: Präsentation vorbereiten',
            'Word: Seriendruck'
        ],
        # Zyklus 6
        [
            'Word: Seriendruck',
            'PowerPoint: Designvorlagen',
            'Word: Dokumentvorlagen',
            'Excel: Power Query Grundlagen',
            'PowerPoint: Präsentation vorbereiten',
            'Word: Umfangreiche Dokumente'
        ],
        # Zyklus 7
        [
            'Word: Dokumentvorlagen',
            'Excel: Power Query Grundlagen',
            'Word: Seriendruck',
            'Word: Umfangreiche Dokumente',
            'PowerPoint: Designvorlagen',
            'PowerPoint: Präsentation vorbereiten'
        ],
        # Zyklus 8
        [
            'Word: Umfangreiche Dokumente',
            'Excel: Power Query Grundlagen',
            'Word: Dokumentvorlagen',
            'Word: Seriendruck',
            'PowerPoint: Präsentation vorbereiten',
            'PowerPoint: Designvorlagen'
        ],
        # Zyklus 9
        [
            'Word: Umfangreiche Dokumente',
            'PowerPoint: Präsentation vorbereiten',
            'Word: Dokumentvorlagen',
            'Excel: Power Query Grundlagen',
            'PowerPoint: Designvorlagen',
            'Word: Seriendruck'
        ],
        # Zyklus 10
        [
            'Word: Seriendruck',
            'PowerPoint: Präsentation vorbereiten',
            'PowerPoint: Designvorlagen',
            'Word: Umfangreiche Dokumente',
            'Excel: Power Query Grundlagen',
            'Word: Dokumentvorlagen'
        ],
        # Zyklus 11
        [
            'Word: Dokumentvorlagen',
            'PowerPoint: Designvorlagen',
            'Word: Seriendruck',
            'Word: Umfangreiche Dokumente',
            'Excel: Power Query Grundlagen',
            'PowerPoint: Präsentation vorbereiten'
        ]
    ]

    # Nicht-lineare Starttage: Di (1), Do (3), Mo (0), Mi (2)
    rot_offsets = [1, 3, 0, 2]

    # Tabellenkopf
    headers = [
        "Datum", "Tag", "KW", "Monat", "Beginn",
        "Prio 1", "Prio 2", "Prio 3",
        "Hinweis", "Freigabe", "Sync_Status", "edoobox_ID"
    ]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )
    monday_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='medium', color='1F4E79'), bottom=Side(style='thin', color='D9D9D9')
    )
    urlaub_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    feiertag_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    start_date = date(2027, 1, 1)
    end_date = date(2027, 12, 31)
    current_date = start_date

    calendar_rows = []
    calendar_weeks = {}
    while current_date <= end_date:
        if current_date.weekday() < 5:
            calendar_rows.append(current_date)
            kw = current_date.isocalendar()[1]
            if current_date.year == 2027:
                if kw not in calendar_weeks:
                    calendar_weeks[kw] = []
                calendar_weeks[kw].append(current_date)
        current_date += timedelta(days=1)

    schedule = {d: {'prio1': '', 'prio2': '', 'prio3': ''} for d in calendar_rows}

    # Nur aktive Unterrichtswochen mit mindestens 3 freien Werktagen
    teaching_kws = [kw for kw in sorted(calendar_weeks.keys()) if len([d for d in calendar_weeks[kw] if not is_blocked(d)]) >= 3]

    for idx, kw in enumerate(teaching_kws):
        cycle_num = idx // 4
        week_in_cycle = idx % 4
        rot_offset = rot_offsets[cycle_num % len(rot_offsets)]
        p3_pool = prio3_rotations[cycle_num % len(prio3_rotations)]

        # Aufbau der Wochenstruktur mit dynamisch zugewiesenem Prio-3-Partner
        if week_in_cycle == 0:
            # Woche 1: Excel GL M1 steht komplett allein
            days_setup = [
                (0, 'Excel Grundlagen Modul 1', None, None),
                (1, 'PowerPoint Grundlagen', None, p3_pool[0]),
                (2, None, 'Word Grundlagen', p3_pool[1])
            ]
            fr_setup = (None, None)
        elif week_in_cycle == 1:
            # Woche 2: Excel GL M2 und Word Aufbau mit Prio 3 parallel
            days_setup = [
                (0, 'Excel Grundlagen Modul 2', None, p3_pool[2]),
                (1, 'PowerPoint Aufbau', None, p3_pool[3]),
                (2, None, 'Word Aufbau', p3_pool[4])
            ]
            fr_setup = (None, None)
        elif week_in_cycle == 2:
            # Woche 3: Outlook GL parallel zu Excel Große Tabellen; Excel GL M3 mit Prio 3
            days_setup = [
                (0, None, 'Excel Grundlagen Modul 3', p3_pool[5]),
                (1, 'Outlook Grundlagen', 'Excel: Große Tabellen', None),
                (2, 'Excel Aufbau Modul 1', None, None)
            ]
            fr_setup = (None, None)
        else:
            # Woche 4: Outlook Aufbau parallel zu Excel Aufbau M2; Excel GL M4 auf Freitag
            days_setup = [
                (0, 'Excel: Pivot-Tabellen', None, None),
                (1, 'Outlook Aufbau', 'Excel Aufbau Modul 2', None)
            ]
            fr_setup = (None, 'Excel Grundlagen Modul 4')

        kw_days = calendar_weeks[kw]
        free_modo = [d for d in kw_days if d.weekday() < 4 and not is_blocked(d)]
        free_fr = [d for d in kw_days if d.weekday() == 4 and not is_blocked(d)]

        used_days_this_week = set()

        for (day_slot, p1_c, p2_c, p3_c) in days_setup:
            target_wd = (day_slot + rot_offset) % 4
            matching = [d for d in free_modo if d.weekday() == target_wd and d not in used_days_this_week]
            if matching:
                assigned_d = matching[0]
            else:
                remaining = [d for d in free_modo if d not in used_days_this_week]
                assigned_d = remaining[0] if remaining else None

            if assigned_d:
                used_days_this_week.add(assigned_d)
                schedule[assigned_d]['prio1'] = p1_c or ''
                schedule[assigned_d]['prio2'] = p2_c or ''
                schedule[assigned_d]['prio3'] = p3_c or ''

        fr_p2, fr_p3 = fr_setup
        if (fr_p2 or fr_p3) and free_fr:
            schedule[free_fr[0]]['prio2'] = fr_p2 or ''
            schedule[free_fr[0]]['prio3'] = fr_p3 or ''

    # 3. Schreiben in Excel mit festen Werten (formelfrei)
    weekdays_de = ["Mo", "Di", "Mi", "Do", "Fr"]
    months_de = [
        "Januar", "Februar", "März", "April", "Mai", "Juni",
        "Juli", "August", "September", "Oktober", "November", "Dezember"
    ]

    row_num = 2
    for d in calendar_rows:
        weekday = d.weekday()
        ws.cell(row=row_num, column=1, value=d)
        ws.cell(row=row_num, column=1).number_format = "yyyy-mm-dd"

        ws.cell(row=row_num, column=2, value=weekdays_de[weekday])
        ws.cell(row=row_num, column=3, value=d.isocalendar()[1])
        ws.cell(row=row_num, column=4, value=months_de[d.month - 1])

        hinweis = ""
        fill_to_apply = None
        if d in feiertage:
            hinweis = f"Feiertag: {feiertage[d]}"
            fill_to_apply = feiertag_fill
        elif d in brueckentage:
            hinweis = brueckentage[d]
            fill_to_apply = feiertag_fill
        elif is_urlaub(d):
            hinweis = "Urlaub"
            fill_to_apply = urlaub_fill

        if hinweis == "":
            ws.cell(row=row_num, column=5, value="09:00")
            ws.cell(row=row_num, column=6, value=schedule[d]['prio1'])
            ws.cell(row=row_num, column=7, value=schedule[d]['prio2'])
            ws.cell(row=row_num, column=8, value=schedule[d]['prio3'])
            ws.cell(row=row_num, column=10, value="Entwurf")
            ws.cell(row=row_num, column=11, value="Offen")
        else:
            ws.cell(row=row_num, column=9, value=hinweis)

        border_to_use = monday_border if weekday == 0 else thin_border
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_num, column=col_idx)
            c.border = border_to_use
            if fill_to_apply and col_idx in [5, 6, 7, 8]:
                c.fill = fill_to_apply
            if col_idx in [1, 2, 3, 4, 5, 10, 11]:
                c.alignment = Alignment(horizontal="center", vertical="center")

        row_num += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 13)

    output_file = "Kursplanung_MS_Office_2027.xlsx"
    wb.save(output_file)
    print(f"Erfolgreich generiert: {output_file}")

if __name__ == "__main__":
    generate_2027_schedule_dynamic_mix()